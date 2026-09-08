"""
把21篇跨主题试点样本的最高分候选文件内容dump成文本,供人工/LLM逐篇判断是否真的是
基因signature+系数表,并做实际抽取。docx文件额外处理(不是openpyxl能读的)。
"""
import json
from pathlib import Path

import openpyxl

PKG_DIR = Path(__file__).resolve().parent.parent / "data" / "pmc_packages"
OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "pilot21_dump"
OUT_DIR.mkdir(exist_ok=True)

PICKS = [
    ("PMC10829884", "mmc2.xlsx"),
    ("PMC11005085", "Supplementary_Data5.xlsx"),
    ("PMC11080313", "12885_2024_12327_MOESM6_ESM.xlsx"),
    ("PMC11302292", "12885_2024_12741_MOESM10_ESM.xlsx"),
    ("PMC11336146", "11357_2024_1164_MOESM1_ESM.xlsx"),
    ("PMC11458410", "Table1.xlsx"),
    ("PMC11462029", "mmc1_extracted/Supplenmentary Tables--revised/Table S3.xlsx"),
    ("PMC11496652", "41598_2024_75650_MOESM2_ESM.xlsx"),
    ("PMC11586327", "12672_2024_1575_MOESM3_ESM.csv"),
    ("PMC11836339", "41598_2025_89770_MOESM3_ESM.xlsx"),
    ("PMC11954311", "13048_2025_1589_MOESM2_ESM.xlsx"),
    ("PMC12047781", "pone.0322618.s005.xlsx"),
    ("PMC12084375", "41598_2025_2134_MOESM1_ESM.xlsx"),
    ("PMC12311075", "12672_2025_3294_MOESM4_ESM.xlsx"),
    ("PMC12567641", "CNR2-8-e70372-s004.xlsx"),
    ("PMC12585956", "Table1.xlsx"),
    ("PMC12702948", "DataSheet1_extracted/Additional file 4.XLSX"),
    ("PMC13138313", "peerj-14-21117-s001.docx"),
    ("PMC13140268", "MI-2026-3900151-s004.xlsx"),
    ("PMC13393646", "mmc2.xlsx"),
    ("PMC13437700", "Table5.xlsx"),
]


def dump_xlsx(path: Path, max_rows=40):
    lines = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        return [f"[ERROR loading: {e}]"]
    for ws in wb.worksheets:
        lines.append(f"--- sheet: {ws.title} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True)):
            cells = [str(c) if c is not None else "" for c in row]
            lines.append(" | ".join(cells))
            if i >= max_rows:
                break
        lines.append("")
    return lines


def main():
    for pmcid, relpath in PICKS:
        fpath = PKG_DIR / pmcid / relpath
        out_txt = OUT_DIR / f"{pmcid}.txt"
        if fpath.suffix.lower() == ".docx":
            lines = ["[docx文件,需要另外用python-docx或textutil读取,这里先跳过]"]
        elif fpath.suffix.lower() == ".csv":
            try:
                lines = fpath.read_text(encoding="utf-8", errors="ignore").splitlines()[:40]
            except Exception as e:
                lines = [f"[ERROR: {e}]"]
        else:
            lines = dump_xlsx(fpath)
        out_txt.write_text(f"=== {pmcid} :: {relpath} ===\n" + "\n".join(lines), encoding="utf-8")
        print(f"{pmcid}: {len(lines)} lines -> {out_txt}")


if __name__ == "__main__":
    main()
